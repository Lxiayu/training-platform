# -*- coding: utf-8 -*-
"""合并设计文档生成 Word 版本。

将《02-总体设计.md》(完整) 与《01-数据库设计.md》(精简版: 省略 DDL/查询/种子 SQL)
合并为 docs/启航编程管理系统-设计文档合集.docx。

用法: python scripts/build_docx.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
DOC_GENERAL = ROOT / "docs" / "02-总体设计.md"
DOC_DB = ROOT / "docs" / "01-数据库设计.md"
OUTPUT = ROOT / "docs" / "启航编程管理系统-设计文档合集.docx"

LINK_RE = re.compile(r"\[([^\]]+)\]\([^)]*\)")
BOLD_RE = re.compile(r"(\*\*.+?\*\*)")
CODE_RE = re.compile(r"(`[^`]+`)")


# ---------- 底层辅助 ----------

def set_font(run, mono=False, bold=None, italic=None, size=None, color=None, east="宋体"):
    """统一设置字体: 西文 + 中文字体分离。size=None 表示继承样式字号。"""
    if mono:
        run.font.name = "Consolas"
        run.font.size = size or Pt(9)
        east = "微软雅黑"
    else:
        run.font.name = "Calibri"
        if size is not None:
            run.font.size = size
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_paragraph(p, fill="F5F5F5"):
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def shade_cell(cell, fill="DEEAF6"):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def add_inline(paragraph, text, bold=False, italic=False, size=None):
    """行内格式渲染: [链接](..) -> 链接文字, **粗体**, `代码`。"""
    text = LINK_RE.sub(r"\1", text)
    for part in BOLD_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            add_inline(paragraph, part[2:-2], bold=True, italic=italic, size=size)
            continue
        for seg in CODE_RE.split(part):
            if not seg:
                continue
            if seg.startswith("`") and seg.endswith("`") and len(seg) > 2:
                r = paragraph.add_run(seg[1:-1])
                set_font(r, mono=True, bold=bold, italic=italic)
            else:
                r = paragraph.add_run(seg)
                set_font(r, bold=bold or None, italic=italic or None, size=size)


# ---------- 块级渲染 ----------

def render_code(document, code_lines):
    for ln in code_lines:
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Cm(0.35)
        shade_paragraph(p)
        r = p.add_run(ln if ln.strip() else " ")
        set_font(r, mono=True)


def render_table(document, rows):
    parsed = []
    for row in rows:
        if re.match(r"^\|[\s:\-|]+\|$", row):  # 跳过分隔行
            continue
        cells = [c.strip() for c in row.strip().strip("|").split("|")]
        parsed.append(cells)
    if not parsed:
        return
    ncols = max(len(r) for r in parsed)
    table = document.add_table(rows=len(parsed), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(parsed):
        for ci in range(ncols):
            text = row[ci] if ci < len(row) else ""
            p = table.cell(ri, ci).paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            add_inline(p, text, bold=(ri == 0), size=Pt(9))
            if ri == 0:
                shade_cell(table.cell(ri, ci))
    tail = document.add_paragraph()
    tail.paragraph_format.space_after = Pt(4)
    tail.paragraph_format.space_before = Pt(0)

# ---------- Markdown 主渲染 ----------

def render_md(document, lines):
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()
        if not s or s == "---" or s == "<!-- APPEND -->":
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            level, text = len(m.group(1)), m.group(2)
            if level == 1:  # 文档主标题由封面替代
                i += 1
                continue
            h = document.add_heading("", level=level)
            add_inline(h, text)
            i += 1
            continue
        if s.startswith("```"):
            i += 1
            code = []
            while i < n and lines[i].strip() != "```":
                code.append(lines[i])
                i += 1
            i += 1
            render_code(document, code)
            continue
        if s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            render_table(document, rows)
            continue
        if s.startswith(">"):
            quote = []
            while i < n and lines[i].strip().startswith(">"):
                q = lines[i].strip()[1:].strip()
                if q:
                    quote.append(q)
                i += 1
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(6)
            shade_paragraph(p, "EFF3FA")
            add_inline(p, " ".join(quote), italic=True, size=Pt(10))
            continue
        m = re.match(r"^(\s*)([-*])\s+(.*)$", line)
        if m:
            indent, _, body = m.groups()
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.1 if len(indent) >= 2 else 0.6)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, "• " + body)
            i += 1
            continue
        m = re.match(r"^(\s*)(\d+)[.、]\s+(.*)$", line)
        if m:
            indent, num, body = m.groups()
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.1 if len(indent) >= 2 else 0.6)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, f"{num}. {body}")
            i += 1
            continue
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        add_inline(p, s)
        i += 1


# ---------- 数据库设计文档精简 ----------

def slim_db_doc(lines):
    """精简: 删目录、mermaid ER 源码、全部 sql 代码块、第 8/9 章(查询与种子 SQL)。"""
    out, i, n = [], 0, len(lines)
    skip_h2, notice_added = False, False
    while i < n:
        line = lines[i]
        s = line.strip()
        if s.startswith("```sql"):
            i += 1
            while i < n and lines[i].strip() != "```":
                i += 1
            i += 1
            continue
        if s.startswith("```mermaid"):
            out.append("> 📌 ER 图 Mermaid 源码从略，关系清单见下方 3.2 节表格。")
            i += 1
            while i < n and lines[i].strip() != "```":
                i += 1
            i += 1
            continue
        if skip_h2:
            if re.match(r"^## ", line) and not re.match(r"^## (目录|8\.|9\.)", line):
                skip_h2 = False
                out.append(line)
            elif not notice_added and re.match(r"^## (9\.|10\.)", line):
                out.append("> 第 8 章（典型查询 SQL）与第 9 章（建库 / 种子数据 SQL）从略，"
                           "完整可执行脚本见《01-数据库设计.md》原文。")
                notice_added = True
            i += 1
            continue
        if re.match(r"^## 目录", line) or re.match(r"^## 8\. ", line):
            skip_h2 = True
            i += 1
            continue
        out.append(line)
        i += 1
    return out

# ---------- 封面 / 目录 / 样式 ----------

def init_document(document):
    for sec in document.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.2)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "宋体")
    for i, sz in [(1, 17), (2, 14), (3, 12), (4, 11)]:
        st = document.styles[f"Heading {i}"]
        st.font.name = "Calibri"
        st.font.size = Pt(sz)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x1F, 0x3B, 0x73)
        srpr = st.element.get_or_add_rPr()
        srf = srpr.find(qn("w:rFonts"))
        if srf is None:
            srf = OxmlElement("w:rFonts")
            srpr.append(srf)
        srf.set(qn("w:eastAsia"), "微软雅黑")


def build_cover(document):
    for _ in range(5):
        document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("启航编程管理系统")
    set_font(r, bold=True, size=Pt(30), east="微软雅黑")
    p2 = document.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("设 计 文 档 合 集")
    set_font(r2, bold=True, size=Pt(22), east="微软雅黑", color=RGBColor(0x1F, 0x3B, 0x73))
    p3 = document.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run("总体设计 · 数据库设计（精简版）")
    set_font(r3, size=Pt(13), color=RGBColor(0x59, 0x59, 0x59))
    for _ in range(4):
        document.add_paragraph()
    info = [
        ("文档版本", "v1.0"),
        ("更新日期", "2026-10-08"),
        ("第一部分", "总体设计（架构 / 模块 / 流程 / 接口概要）"),
        ("第二部分", "数据库设计 · 精简版（模型 / 枚举 / 表结构 / 业务规则）"),
        ("源文档", "docs/02-总体设计.md · docs/01-数据库设计.md"),
        ("文档状态", "初稿 · 待评审"),
    ]
    t = document.add_table(rows=len(info), cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, (k, v) in enumerate(info):
        p = t.cell(ri, 0).paragraphs[0]
        add_inline(p, k, bold=True, size=Pt(10.5))
        shade_cell(t.cell(ri, 0))
        p = t.cell(ri, 1).paragraphs[0]
        add_inline(p, v, size=Pt(10.5))


def add_toc_placeholder(document):
    document.add_heading("目 录", level=1)
    p = document.add_paragraph()
    parts = []
    for kind, value in (("begin", None), ("instr", ' TOC \\o "1-3" \\h \\z \\u '),
                        ("separate", None), ("text",
                        "（请在 Word 中右键此处 → 更新域，或全选后按 F9 生成目录）"),
                        ("end", None)):
        r = p.add_run()
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = value
        elif kind == "text":
            el = OxmlElement("w:t")
            el.text = value
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        parts.append((r, el))
    for r, el in parts:
        r._r.append(el)


# ---------- 主流程 ----------

def main():
    document = Document()
    init_document(document)
    build_cover(document)
    document.add_page_break()
    add_toc_placeholder(document)
    document.add_page_break()

    document.add_heading("第一部分 总体设计", level=1)
    render_md(document, DOC_GENERAL.read_text(encoding="utf-8").splitlines())
    document.add_page_break()

    document.add_heading("第二部分 数据库设计（精简版）", level=1)
    note = document.add_paragraph()
    note.paragraph_format.left_indent = Cm(0.5)
    shade_paragraph(note, "EFF3FA")
    add_inline(note, "本部分为《01-数据库设计.md》v1.0 的精简版：保留设计约定、关系模型、"
                     "枚举字典、各表字段定义与业务规则；省略全部可执行 DDL、典型查询 SQL 与"
                     "种子数据脚本，完整内容请查阅源文档。", italic=True, size=Pt(10))
    render_md(document, slim_db_doc(DOC_DB.read_text(encoding="utf-8").splitlines()))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f"[OK] generated: {OUTPUT}")
    print(f"     size: {OUTPUT.stat().st_size} bytes")


if __name__ == "__main__":
    main()


